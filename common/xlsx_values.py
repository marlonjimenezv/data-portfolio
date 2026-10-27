"""Recalculate a workbook in LibreOffice and read back the formula results."""
from __future__ import annotations

import datetime as dt
import subprocess
import warnings
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

COMMON = Path(__file__).resolve().parent
EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def lo_recalc(src: Path, out_values: Path, pdf: Path | None = None, sheets: list[str] | None = None,
              one_page: list[str] | None = None, areas: dict | None = None) -> None:
    """Run common/lo_tools.py with the system python. Raises if any formula returns an error."""
    cmd = ["/usr/bin/python3", str(COMMON / "lo_tools.py"), "recalc", str(src), str(out_values)]
    if pdf:
        cmd += ["--pdf", str(pdf)]
    if sheets:
        cmd += ["--sheets", ",".join(sheets)]
    if one_page:
        cmd += ["--one-page", ",".join(one_page)]
    if areas:
        cmd += ["--areas", ";".join(f"{k}={v}" for k, v in areas.items())]
    res = subprocess.run(cmd, capture_output=True, text=True)
    print(res.stdout.rstrip())
    if res.returncode != 0:
        raise SystemExit(f"LibreOffice found formula errors in {src.name}:\n{res.stdout}\n{res.stderr}")


def read_values(path: Path, positions: list[tuple[str, int, int]]) -> dict:
    """Return {(sheet, row0, col0): value} for the given 0-based positions.

    Dates are converted to Excel serial numbers so they can be stored as cached values.
    """
    wanted = defaultdict(set)
    for sh, r, c in positions:
        wanted[sh].add((r, c))
    out = {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, data_only=True, read_only=True)
    for sh, cells in wanted.items():
        ws = wb[sh]
        max_r = max(r for r, _ in cells) + 1
        max_c = max(c for _, c in cells) + 1
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=True)):
            for c_idx, v in enumerate(row):
                if (r_idx, c_idx) in cells:
                    if isinstance(v, dt.datetime):
                        v = (v - EXCEL_EPOCH).total_seconds() / 86400
                    elif isinstance(v, dt.date):
                        v = (dt.datetime.combine(v, dt.time()) - EXCEL_EPOCH).days
                    out[(sh, r_idx, c_idx)] = "" if v is None else v
    wb.close()
    return out
